
import calendar
import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.db.models import Count, Q, Sum
from django.db.models.functions import ExtractHour, ExtractMonth
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import User, Customer, Staff
from appointments.models import Appointment, Availability


def status_counts(appointments):
    return {
        "total": appointments.count(),
        "completed": appointments.filter(
            status=Appointment.Status.COMPLETED
        ).count(),
        "scheduled": appointments.filter(
            status=Appointment.Status.SCHEDULED
        ).count(),
        "cancelled": appointments.filter(
            status=Appointment.Status.CANCELLED
        ).count(),
        "rescheduled": appointments.filter(
            status=Appointment.Status.RESCHEDULED
        ).count(),
    }


def appointment_history_row(appointment):
    customer = appointment.customer
    staff = appointment.staff
    service = appointment.service

    return {
        "appointment_id": appointment.id,
        "appointment_date": appointment.appointment_date,
        "start_time": appointment.start_time,
        "end_time": appointment.end_time,
        "status": appointment.status,
        "customer_id": appointment.customer_id,
        "customer_name": (
            f"{customer.user.first_name} {customer.user.last_name}"
        ).strip(),
        "customer_email": customer.user.email,
        "customer_phone": customer.phone,
        "staff_id": appointment.staff_id,
        "staff_name": (
            f"{staff.user.first_name} {staff.user.last_name}"
        ).strip(),
        "staff_type": staff.staff_type,
        "service_id": appointment.service_id,
        "service_name": service.name,
        "service_price": service.price,
        "reason": appointment.reason,
        "cancellation_reason": appointment.cancellation_reason,
        "created_at": appointment.created_at,
        "updated_at": appointment.updated_at,
    }


class AdminReportBaseView(APIView):
    permission_classes = [IsAuthenticated]

    def check_admin(self, request):
        if not request.user.is_authenticated:
            return Response(
                {"detail": "Authentication is required."},
                status=401,
            )

        if request.user.role != User.Role.ADMIN:
            return Response(
                {"detail": "Only administrators can access reports."},
                status=403,
            )

        return None

    def get_date_range(self, request):
        from_text = request.query_params.get("from")
        to_text = request.query_params.get("to")

        from_date = parse_date(from_text) if from_text else None
        to_date = parse_date(to_text) if to_text else None

        if from_text and not from_date:
            raise ValueError(
                "'from' must be a valid date in YYYY-MM-DD format."
            )

        if to_text and not to_date:
            raise ValueError(
                "'to' must be a valid date in YYYY-MM-DD format."
            )

        if from_date and not to_date:
            to_date = from_date

        if to_date and not from_date:
            from_date = to_date

        if from_date and to_date and from_date > to_date:
            raise ValueError("'from' date cannot be after 'to' date.")

        return from_date, to_date

    def filtered_appointments(self, request):
        appointments = Appointment.objects.select_related(
            "customer__user",
            "staff__user",
            "service",
        )

        from_date, to_date = self.get_date_range(request)

        if from_date:
            appointments = appointments.filter(
                appointment_date__gte=from_date
            )

        if to_date:
            appointments = appointments.filter(
                appointment_date__lte=to_date
            )

        return appointments, from_date, to_date

    def get_report_data(self, request):
        try:
            appointments, from_date, to_date = (
                self.filtered_appointments(request)
            )
        except ValueError as exc:
            return None, None, None, Response(
                {"detail": str(exc)},
                status=400,
            )

        return appointments, from_date, to_date, None

    def excel_value(self, value):
        if isinstance(value, Decimal):
            return float(value)

        if isinstance(value, dict):
            return json.dumps(value, default=str, ensure_ascii=False)

        if isinstance(value, (list, tuple)):
            return json.dumps(value, default=str, ensure_ascii=False)

        if isinstance(value, datetime) and timezone.is_aware(value):
            return timezone.localtime(value).replace(tzinfo=None)

        if hasattr(value, "isoformat") and not isinstance(value, str):
            return value

        return value

    def add_table_sheet(self, workbook, title, rows):
        sheet = workbook.create_sheet(title=title[:31])

        if isinstance(rows, dict):
            rows = [rows]

        if not rows:
            sheet.append(["No records found"])
            sheet.column_dimensions["A"].width = 24
            return

        if not isinstance(rows, list):
            rows = [{"value": rows}]

        if not all(isinstance(row, dict) for row in rows):
            rows = [{"value": item} for item in rows]

        headers = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)

        sheet.append(headers)

        for row in rows:
            sheet.append([
                self.excel_value(row.get(header))
                for header in headers
            ])

        for column_cells in sheet.columns:
            column_letter = get_column_letter(
                column_cells[0].column
            )
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                45,
            )

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    def export_excel(self, response, request):
        if response.status_code >= 400:
            return response

        data = response.data
        workbook = Workbook()
        workbook.remove(workbook.active)

        if isinstance(data, dict):
            summary = {}
            tables = []

            for key, value in data.items():
                if isinstance(value, list):
                    if not value:
                        tables.append((key, []))
                    elif all(isinstance(item, dict) for item in value):
                        tables.append((key, value))
                    else:
                        summary[key] = json.dumps(
                            value,
                            default=str,
                            ensure_ascii=False,
                        )
                elif isinstance(value, dict):
                    for nested_key, nested_value in value.items():
                        summary[f"{key}_{nested_key}"] = nested_value
                else:
                    summary[key] = value

            if summary:
                self.add_table_sheet(
                    workbook,
                    "Summary",
                    [summary],
                )

               
            for title, rows in tables:
                if (
                    self.__class__.__name__ == "MonthlyAppointmentReportView"
                    and title == "data"
                ):
                    title = "Monthly Summary"

                self.add_table_sheet(workbook, title, rows)

            if self.__class__.__name__ == "MonthlyAppointmentReportView":
                details = self.get_export_details(request)

                self.add_table_sheet(
                    workbook,
                    "Appointment Details",
                    details,
                )


        elif isinstance(data, list):
            self.add_table_sheet(workbook, "Report", data)
        else:
            self.add_table_sheet(
                workbook,
                "Report",
                [{"value": data}],
            )

        if not workbook.sheetnames:
            self.add_table_sheet(workbook, "Report", [])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        export_time = timezone.localtime(timezone.now())
        date_part = export_time.strftime("%d_%B_%Y")
        time_part = export_time.strftime("%I-%M-%p")

        report_name = (
            self.__class__.__name__
            .replace("ReportView", "")
            .replace("View", "")
        )

        filename = (
            f"{date_part}_{time_part}_{report_name}.xlsx"
        )

        excel_response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        excel_response["Content-Disposition"] = (
            f'attachment; filename="{filename}"'
        )
        return excel_response

    def finalize_response(
        self,
        request,
        response,
        *args,
        **kwargs,
    ):
        response = super().finalize_response(
            request,
            response,
            *args,
            **kwargs,
        )

        if (
            request.query_params.get("format", "").lower() == "xlsx"
            and isinstance(response, Response)
            and response.status_code < 400
        ):
            return self.export_excel(response, request)

        return response


class AdminDashboardReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        counts = status_counts(appointments)

        completed = appointments.filter(
            status=Appointment.Status.COMPLETED
        )

        estimated_value = (
            completed.aggregate(total=Sum("service__price"))["total"]
            or Decimal("0.00")
        )

        customers = User.objects.filter(role=User.Role.CUSTOMER)

        if from_date:
            customers = customers.filter(
                created_at__date__gte=from_date
            )

        if to_date:
            customers = customers.filter(
                created_at__date__lte=to_date
            )

        total = counts["total"]

        return Response({
            "period": {
                "from": from_date,
                "to": to_date,
            },
            "kpis": {
                "total_appointments": total,
                "completed": counts["completed"],
                "scheduled": counts["scheduled"],
                "cancelled": counts["cancelled"],
                "rescheduled": counts["rescheduled"],
                "completion_rate": (
                    round(counts["completed"] / total * 100, 2)
                    if total else 0
                ),
                "cancellation_rate": (
                    round(counts["cancelled"] / total * 100, 2)
                    if total else 0
                ),
                "estimated_completed_service_value": estimated_value,
                "new_customers": customers.count(),
            },
        })


class DailyAppointmentReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        if not from_date or not to_date:
            return Response(
                {"detail": "'from' and 'to' query parameters are required."},
                status=400,
            )

        rows = []
        for day in (
            appointments.values_list("appointment_date", flat=True)
            .distinct()
            .order_by("appointment_date")
        ):
            daily = appointments.filter(appointment_date=day)
            counts = status_counts(daily)
            rows.append({
                "date": day,
                "total": counts["total"],
                "completed": counts["completed"],
                "scheduled": counts["scheduled"],
                "cancelled": counts["cancelled"],
                "rescheduled": counts["rescheduled"],
            })

        return Response({
            "period": {"from": from_date, "to": to_date},
            "data": rows,
        })




class MonthlyAppointmentReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        year_text = request.query_params.get("year")

        if not year_text:
            return Response(
                {"detail": "'year' query parameter is required."},
                status=400,
            )

        try:
            year = int(year_text)

            if year < 2000 or year > 2100:
                raise ValueError

        except (ValueError, TypeError):
            return Response(
                {"detail": "'year' must be a valid year."},
                status=400,
            )

        appointments = Appointment.objects.filter(
            appointment_date__year=year
        )

        rows = []

        for month_number in range(1, 13):
            monthly = appointments.filter(
                appointment_date__month=month_number
            )

            counts = status_counts(monthly)

            rows.append({
                "month": calendar.month_name[month_number],
                "month_number": month_number,
                "total": counts["total"],
                "completed": counts["completed"],
                "scheduled": counts["scheduled"],
                "cancelled": counts["cancelled"],
                "rescheduled": counts["rescheduled"],
            })

        return Response({
            "year": year,
            "data": rows,
        })

    def get_export_sheet_title(self, key):
        if key == "data":
            return "Monthly Summary"

        return key

    def get_export_details(self, request):
        year = int(request.query_params["year"])

        appointments = (
            Appointment.objects
            .filter(appointment_date__year=year)
            .select_related(
                "customer",
                "staff",
                "service",
            )
            .order_by(
                "appointment_date",
                "start_time",
            )
        )

        details = []

        for appointment in appointments:
            customer = appointment.customer
            staff = appointment.staff

            customer_name = (
                f"{customer.user.first_name} "
                f"{customer.user.last_name}"
            ).strip()

            doctor_name = (
                f"{staff.user.first_name} "
                f"{staff.user.last_name}"
            ).strip()

            details.append({
                "Appointment ID": appointment.id,
                "Customer Name": customer_name,
                "Customer Email": customer.user.email,
                "Doctor Name": doctor_name,
                "Service": str(appointment.service),
                "Appointment Date": appointment.appointment_date,
                "Start Time": appointment.start_time,
                "End Time": appointment.end_time,
                "Status": appointment.get_status_display(),
                "Booking Date": appointment.created_at,
                "Last Updated": appointment.updated_at,
                "Cancellation Reason": (
                    appointment.cancellation_reason
                    if appointment.status == Appointment.Status.CANCELLED
                    else ""
                ),
                "Appointment Reason": appointment.reason or "",
            })

        return details

class DateRangeAppointmentReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        if not from_date or not to_date:
            return Response(
                {"detail": "'from' and 'to' query parameters are required."},
                status=400,
            )

        counts = status_counts(appointments)
        total = counts["total"]

        return Response({
            "period": {"from": from_date, "to": to_date},
            "total_appointments": total,
            "completed": counts["completed"],
            "scheduled": counts["scheduled"],
            "cancelled": counts["cancelled"],
            "rescheduled": counts["rescheduled"],
            "completion_rate": (
                round(counts["completed"] / total * 100, 2)
                if total else 0
            ),
            "cancellation_rate": (
                round(counts["cancelled"] / total * 100, 2)
                if total else 0
            ),
        })


class DoctorPerformanceReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        rows = (
            appointments
            .filter(staff__staff_type="DOCTOR")
            .values(
                "staff_id",
                "staff__user__first_name",
                "staff__user__last_name",
            )
            .order_by("staff_id")
        )

        data = []
        for row in rows:
            doctor_appointments = appointments.filter(
                staff_id=row["staff_id"]
            )
            counts = status_counts(doctor_appointments)
            total = counts["total"]

            data.append({
                "doctor_id": row["staff_id"],
                "doctor_name": (
                    f'{row["staff__user__first_name"]} '
                    f'{row["staff__user__last_name"]}'
                ).strip(),
                "total_appointments": total,
                "completed": counts["completed"],
                "scheduled": counts["scheduled"],
                "cancelled": counts["cancelled"],
                "rescheduled": counts["rescheduled"],
                "completion_rate": (
                    round(counts["completed"] / total * 100, 2)
                    if total else 0
                ),
            })

        data.sort(
            key=lambda item: item["total_appointments"],
            reverse=True,
        )

        return Response({
            "period": {"from": from_date, "to": to_date},
            "data": data,
        })


class ServiceUtilizationReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        rows = (
            appointments
            .values("service_id", "service__name")
            .order_by("service_id")
        )

        data = []
        for row in rows:
            service_appointments = appointments.filter(
                service_id=row["service_id"]
            )
            completed = service_appointments.filter(
                status=Appointment.Status.COMPLETED
            )

            data.append({
                "service_id": row["service_id"],
                "service_name": row["service__name"],
                "appointments": service_appointments.count(),
                "completed": completed.count(),
                "cancelled": service_appointments.filter(
                    status=Appointment.Status.CANCELLED
                ).count(),
                "estimated_completed_service_value": (
                    completed.aggregate(
                        total=Sum("service__price")
                    )["total"] or Decimal("0.00")
                ),
            })

        data.sort(
            key=lambda item: item["appointments"],
            reverse=True,
        )

        return Response({
            "period": {"from": from_date, "to": to_date},
            "data": data,
        })


class CancellationReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        cancelled = appointments.filter(
            status=Appointment.Status.CANCELLED
        )

        total = appointments.count()
        cancelled_count = cancelled.count()

        by_date = (
            cancelled.values("appointment_date")
            .annotate(total=Count("id"))
            .order_by("appointment_date")
        )

        by_staff = (
            cancelled.values(
                "staff_id",
                "staff__user__first_name",
                "staff__user__last_name",
            )
            .annotate(total=Count("id"))
            .order_by("-total")
        )

        return Response({
            "period": {"from": from_date, "to": to_date},
            "total_cancelled": cancelled_count,
            "cancellation_rate": (
                round(cancelled_count / total * 100, 2)
                if total else 0
            ),
            "by_date": [
                {
                    "date": row["appointment_date"],
                    "cancelled": row["total"],
                }
                for row in by_date
            ],
            "by_staff": [
                {
                    "staff_id": row["staff_id"],
                    "staff_name": (
                        f'{row["staff__user__first_name"]} '
                        f'{row["staff__user__last_name"]}'
                    ).strip(),
                    "cancelled": row["total"],
                }
                for row in by_staff
            ],
            "cancellation_details": [
                appointment_history_row(appointment)
                for appointment in cancelled.order_by(
                    "appointment_date",
                    "start_time",
                )
            ],
        })


class PeakHoursReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        date_text = request.query_params.get("date")
        if not date_text:
            return Response(
                {"detail": "'date' query parameter is required."},
                status=400,
            )

        selected_date = parse_date(date_text)
        if not selected_date:
            return Response(
                {"detail": "'date' must be in YYYY-MM-DD format."},
                status=400,
            )

        appointments = (
            Appointment.objects
            .filter(appointment_date=selected_date)
            .exclude(status=Appointment.Status.CANCELLED)
        )

        rows = (
            appointments
            .annotate(hour=ExtractHour("start_time"))
            .values("hour")
            .annotate(appointments=Count("id"))
            .order_by("-appointments", "hour")
        )

        return Response({
            "date": selected_date,
            "peak_hours": [
                {
                    "hour": (
                        f'{row["hour"]:02d}:00'
                        if row["hour"] is not None
                        else None
                    ),
                    "appointments": row["appointments"],
                }
                for row in rows
            ],
        })


class RevenueReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        completed = appointments.filter(
            status=Appointment.Status.COMPLETED
        )

        summary = completed.aggregate(
            estimated_service_value=Sum("service__price"),
            completed_appointments=Count("id"),
        )

        estimated_value = (
            summary["estimated_service_value"] or Decimal("0.00")
        )
        completed_count = summary["completed_appointments"] or 0

        service_rows = (
            completed
            .values("service_id", "service__name")
            .order_by("service_id")
        )

        service_breakdown = []
        for row in service_rows:
            service_appointments = completed.filter(
                service_id=row["service_id"]
            )
            service_breakdown.append({
                "service_id": row["service_id"],
                "service_name": row["service__name"],
                "appointments": service_appointments.count(),
                "estimated_service_value": (
                    service_appointments.aggregate(
                        total=Sum("service__price")
                    )["total"] or Decimal("0.00")
                ),
            })

        service_breakdown.sort(
            key=lambda item: item["estimated_service_value"],
            reverse=True,
        )

        return Response({
            "period": {"from": from_date, "to": to_date},
            "completed_appointments": completed_count,
            "estimated_completed_service_value": estimated_value,
            "average_service_value_per_completed_appointment": (
                round(float(estimated_value) / completed_count, 2)
                if completed_count else 0
            ),
            "note": (
                "This is an estimated service value, not verified "
                "payment revenue. The current models do not store payments."
            ),
            "service_breakdown": service_breakdown,
        })


class CustomerAppointmentHistoryReportView(AdminReportBaseView):
    def get(self, request, customer_id):
        denied = self.check_admin(request)
        if denied:
            return denied

        try:
            customer = Customer.objects.select_related("user").get(
                pk=customer_id
            )
        except Customer.DoesNotExist:
            return Response(
                {"detail": "Customer not found."},
                status=404,
            )

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        appointments = appointments.filter(
            customer=customer
        ).order_by("appointment_date", "start_time")

        return Response({
            "customer": {
                "customer_id": customer.id,
                "name": (
                    f"{customer.user.first_name} "
                    f"{customer.user.last_name}"
                ).strip(),
                "email": customer.user.email,
                "phone": customer.phone,
            },
            "period": {"from": from_date, "to": to_date},
            "total_appointments": appointments.count(),
            "history": [
                appointment_history_row(appointment)
                for appointment in appointments
            ],
        })


class DoctorAssignedAppointmentHistoryReportView(AdminReportBaseView):
    def get(self, request, doctor_id):
        denied = self.check_admin(request)
        if denied:
            return denied

        try:
            doctor = Staff.objects.select_related("user").get(
                pk=doctor_id,
                staff_type="DOCTOR",
            )
        except Staff.DoesNotExist:
            return Response(
                {"detail": "Doctor not found."},
                status=404,
            )

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        appointments = appointments.filter(
            staff=doctor
        ).order_by("appointment_date", "start_time")

        return Response({
            "doctor": {
                "staff_id": doctor.id,
                "name": (
                    f"{doctor.user.first_name} "
                    f"{doctor.user.last_name}"
                ).strip(),
                "email": doctor.user.email,
                "phone": doctor.phone,
                "designation": doctor.designation,
                "department": doctor.department,
            },
            "period": {"from": from_date, "to": to_date},
            "total_appointments": appointments.count(),
            "history": [
                appointment_history_row(appointment)
                for appointment in appointments
            ],
        })


class StaffAppointmentHistoryReportView(AdminReportBaseView):
    def get(self, request, staff_id):
        denied = self.check_admin(request)
        if denied:
            return denied

        try:
            staff = Staff.objects.select_related("user").get(
                pk=staff_id
            )
        except Staff.DoesNotExist:
            return Response(
                {"detail": "Staff member not found."},
                status=404,
            )

        appointments, from_date, to_date, error = (
            self.get_report_data(request)
        )
        if error:
            return error

        appointments = appointments.filter(
            staff=staff
        ).order_by("appointment_date", "start_time")

        return Response({
            "staff": {
                "staff_id": staff.id,
                "name": (
                    f"{staff.user.first_name} "
                    f"{staff.user.last_name}"
                ).strip(),
                "email": staff.user.email,
                "phone": staff.phone,
                "staff_type": staff.staff_type,
                "designation": staff.designation,
                "department": staff.department,
                "is_available": staff.is_available,
            },
            "period": {"from": from_date, "to": to_date},
            "total_appointments": appointments.count(),
            "history": [
                appointment_history_row(appointment)
                for appointment in appointments
            ],
        })


class StaffAvailabilityReportView(AdminReportBaseView):
    def get(self, request):
        denied = self.check_admin(request)
        if denied:
            return denied

        availabilities = Availability.objects.select_related(
            "staff__user"
        )

        staff_id = request.query_params.get("staff_id")
        if staff_id:
            try:
                staff_id = int(staff_id)
            except ValueError:
                return Response(
                    {"detail": "'staff_id' must be an integer."},
                    status=400,
                )

            availabilities = availabilities.filter(staff_id=staff_id)

        day_text = request.query_params.get("day_of_week")
        if day_text is not None:
            try:
                day = int(day_text)
                if day < 0 or day > 6:
                    raise ValueError
            except ValueError:
                return Response(
                    {
                        "detail": (
                            "'day_of_week' must be between 0 and 6, "
                            "where Monday is 0 and Sunday is 6."
                        )
                    },
                    status=400,
                )

            availabilities = availabilities.filter(day_of_week=day)

        availabilities = availabilities.order_by(
            "staff__user__first_name",
            "staff__user__last_name",
            "day_of_week",
            "start_time",
        )

        data = [
            {
                "availability_id": item.id,
                "staff_id": item.staff_id,
                "staff_name": (
                    f"{item.staff.user.first_name} "
                    f"{item.staff.user.last_name}"
                ).strip(),
                "staff_type": item.staff.staff_type,
                "day_of_week": item.day_of_week,
                "day_name": item.get_day_of_week_display(),
                "start_time": item.start_time,
                "end_time": item.end_time,
                "is_available": item.is_available,
                "created_at": item.created_at,
                "updated_at": item.updated_at,
            }
            for item in availabilities
        ]

        return Response({
            "total_availability_records": len(data),
            "data": data,
            "note": (
                "This report shows saved availability records. "
                "It does not reconstruct previous schedule changes."
            ),
        })
